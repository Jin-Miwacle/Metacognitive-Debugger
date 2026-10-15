"""
Pulls data out of Supabase (and the AI backend) into spreadsheets your
group can score by hand -- the first step of the hint/quiz evaluation
plan.

Needs two new lines in backend/.env, alongside the Gemini/model ones
already there:

    SUPABASE_URL=https://YOUR-PROJECT-REF.supabase.co
    SUPABASE_SERVICE_KEY=your-secret-service-role-key

Get SUPABASE_SERVICE_KEY from Supabase: Settings -> API -> "service_role".
That's NOT the anon key already in config.js -- it's a separate, secret
key that skips Row Level Security entirely, which is exactly why this
script needs it (it has to read every student's data, not just one
person's). Because of that: never put this key in config.js, never
commit it, never share it. Keep it in backend/.env only, on your own
computer, and only use it for scripts like this one.

Usage (run from the backend/ folder):
    python export_for_evaluation.py hints   -> hints_for_rating.xlsx
    python export_for_evaluation.py quiz    -> quiz_for_rating.xlsx

"hints" reads every saved debug thread straight from the database.
"quiz" needs the AI backend running locally (python main.py, or however
you normally start it) -- it asks /quiz for a fresh quiz on every
reading currently in the database, since individual quiz questions
aren't saved anywhere until a student actually answers one.
"""
import os
import sys

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_KEY")
API_BASE = os.getenv("API_BASE", "http://localhost:8000")

HEADER_FILL = PatternFill(start_color="1F6F5C", end_color="1F6F5C", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10.5)


def supa_get(table, params=""):
    if not SUPABASE_URL or not SUPABASE_SERVICE_KEY:
        sys.exit("Set SUPABASE_URL and SUPABASE_SERVICE_KEY in backend/.env first -- see the top of this file.")
    url = f"{SUPABASE_URL}/rest/v1/{table}?{params}"
    headers = {"apikey": SUPABASE_SERVICE_KEY, "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}"}
    r = requests.get(url, headers=headers, timeout=30)
    r.raise_for_status()
    return r.json()


def style_header(ws, headers, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_row(ws, wrap_cols, total_cols):
    r = ws.max_row
    for c in range(1, total_cols + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = BODY_FONT
        if c in wrap_cols:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_hints():
    print("Fetching debug threads from Supabase...")
    events = supa_get("debug_events", "select=id,session_id,passage,problem,status,transcript&order=id")
    readings = {r["id"]: r for r in supa_get("readings", "select=id,title")}
    sessions = {s["id"]: s for s in supa_get("sessions", "select=id,reading_id")}

    wb = Workbook()
    ws = wb.active
    ws.title = "Hints"
    headers = [
        "Thread ID", "Reading", "Passage", "Problem type", "Status", "Transcript",
        "One move only (not the full answer)?", "Stayed under ~60 words?",
        "Didn't repeat an earlier hint?", "Matched the problem type?",
        "Rater", "Notes",
    ]
    style_header(ws, headers, [10, 22, 34, 16, 12, 60, 14, 14, 14, 14, 12, 30])

    skipped = 0
    for e in events:
        transcript = e.get("transcript") or []
        if not transcript:
            skipped += 1
            continue
        session = sessions.get(e["session_id"], {})
        reading = readings.get(session.get("reading_id"), {})
        lines = []
        for m in transcript:
            who = "AI" if m.get("from") == "ai" else "Student"
            tag = " (offline demo -- not real model output)" if m.get("demo") else ""
            lines.append(f"{who}{tag}: {m.get('text', '')}")
        ws.append([
            e["id"], reading.get("title", "(unknown)"), e.get("passage", ""),
            e.get("problem", ""), e.get("status", ""), "\n".join(lines),
            "", "", "", "", "", "",
        ])
        style_row(ws, wrap_cols={3, 6}, total_cols=len(headers))

    out = "hints_for_rating.xlsx"
    wb.save(out)
    print(f"Wrote {ws.max_row - 1} debug threads to {out} ({skipped} skipped -- no hint saved for those yet).")


def export_quiz(n_per_reading=4):
    print("Fetching readings from Supabase...")
    readings = supa_get("readings", "select=id,title,body&order=created_at")
    if not readings:
        sys.exit("No readings found in the database yet.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Quiz"
    headers = [
        "Reading", "Question", "Kind", "Option A", "Option B", "Option C", "Option D",
        "Model's stated answer", "Model's reasoning ('why')", "Evidence quoted",
        "Stated answer is actually correct?", "All 3 distractors clearly wrong?", "Notes",
    ]
    style_header(ws, headers, [20, 32, 10, 18, 18, 18, 18, 16, 28, 24, 16, 16, 30])

    for r in readings:
        print(f"  generating quiz for: {r['title']}")
        try:
            resp = requests.post(f"{API_BASE}/quiz", json={"text": r["body"], "n": n_per_reading}, timeout=120)
            resp.raise_for_status()
            questions = resp.json().get("questions", [])
        except Exception as e:
            print(f"    skipped ({e})")
            continue
        for q in questions:
            opts = (q.get("options") or []) + ["", "", "", ""]
            idx = q.get("answer")
            stated = opts[idx] if isinstance(idx, int) and 0 <= idx < 4 else ""
            ws.append([
                r["title"], q.get("q", ""), q.get("kind", ""), opts[0], opts[1], opts[2], opts[3],
                stated, q.get("why", ""), q.get("evidence", ""), "", "", "",
            ])
            style_row(ws, wrap_cols={2, 9, 10}, total_cols=len(headers))

    out = "quiz_for_rating.xlsx"
    wb.save(out)
    print(f"Wrote {ws.max_row - 1} quiz questions to {out}.")


if __name__ == "__main__":
    if len(sys.argv) != 2 or sys.argv[1] not in ("hints", "quiz"):
        sys.exit("Usage: python export_for_evaluation.py hints|quiz")
    (export_hints if sys.argv[1] == "hints" else export_quiz)()
